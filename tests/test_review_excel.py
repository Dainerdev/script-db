"""T13: escribir y leer el Excel de revisión (RF-38, RF-39, RF-42 a RF-47, RF-75 a RF-77, RF-92, RF-95, RF-96)."""
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src import parameters as p
from src import review_excel as rx

NAMES = "NOMBRES_APELLIDOS"
E = p.TableEntry


def _entries():
    return [
        E("SOLICITUD", "Devoución", "DEVOLUCIÓN", 95.2, "error de tipeo", p.AUTOMATIC, [1, 2, 3]),
        E("DECISIÓN", "agendada", "AGENDADO", 87.5, "similitud 80–89", p.REVIEW, [4]),
        E("DECISIÓN", "agendadas", "AGENDADO", 82.0, "similitud 80–89", p.REVIEW, [5, 6, 7, 8]),
        E("SALA", "SRVR", "SRVR", None, "confirmar código de SALA", p.TO_CONFIRM, [1, 2]),
        E("TIPO DE EXPEDIENTE", "44317", "POR DETERMINAR", None, "otra categoría", p.AUTOMATIC, [9]),
        E("IDENTIFICACIÓN", "0012", "", None, "identificación inválida", p.REVIEW, [10], decision=None),
        E(NAMES, "Ana Ruiz\nLuis Mora", "Ana Ruiz | Luis Mora", None, "varios nombres", p.DECIDED, [11],
          decision="Ana Ruiz | Luis Mora"),
        E("SOLICITUD", "Rewquerimiento", "Rewquerimiento", None, "decisión", p.DECIDED, [12], decision="MANTENER"),
    ]


def _data(**kwargs):
    base = dict(
        entries=_entries(),
        summary={"DECISIÓN": {"estados": {p.SUMMARY_UNCHANGED: (1, 6), p.REVIEW: (2, 5)}, "otra categoría": (0, 0)}},
        totals={"DECISIÓN": Counter({p.UNCHANGED: 6, p.PENDING_RESULT: 5})},
        catalogs={"SOLICITUD": pd.DataFrame({"Valor": ["DEVOLUCIÓN", "POR DETERMINAR"], "Filas": [7, 1]})},
        separation={"SEPARADA": 3, "SEPARADA SIN CONFIRMAR": 2, "SIN SEPARAR": 1, "": 10, "filas": 25},
        unmatched_decisions=2,
        invalid_decisions=["TIPO DE EXPEDIENTE · «Digital»: «DIGITAL» no está en la lista de la columna; no se aplicó."],
        saved_as="resultado_revision.xlsx",
    )
    base.update(kwargs)
    return rx.ReviewData(**base)


@pytest.fixture
def ruta(tmp_path):
    destino = Path(tmp_path) / "resultado_revision.xlsx"
    rx.write_review(destino, _data())
    return destino


def _hoja(ruta, nombre):
    return pd.read_excel(ruta, sheet_name=nombre, dtype=object, engine="calamine").fillna("")


def _escribir_celda(ruta, hoja, fila_por, valor_por, columna, valor):
    libro = openpyxl.load_workbook(ruta)
    hoja_xl = libro[hoja]
    encabezados = [celda.value for celda in hoja_xl[1]]
    col_busca, col_pone = encabezados.index(fila_por) + 1, encabezados.index(columna) + 1
    for fila in range(2, hoja_xl.max_row + 1):
        if hoja_xl.cell(fila, col_busca).value == valor_por:
            hoja_xl.cell(fila, col_pone).value = valor
    libro.save(ruta)


# ---------------------------------------------------------------------
# Tabla de estandarización: orden, texto, decisiones (RF-38, RF-39, RF-46)
# ---------------------------------------------------------------------

def test_tabla_con_sus_columnas_y_primero_lo_pendiente_por_filas(ruta):
    tabla = _hoja(ruta, "Tabla de estandarización")

    assert list(tabla.columns) == rx.TABLE_COLUMNS
    assert list(tabla["Valor original"][:4]) == ["agendadas", "agendada", "0012", "SRVR"]
    assert list(tabla["Estado"][:4]) == [p.REVIEW, p.REVIEW, p.REVIEW, p.TO_CONFIRM]
    assert list(tabla["Filas afectadas"][:2]) == ["4", "1"]


def test_todo_se_escribe_como_texto(ruta):
    hoja = openpyxl.load_workbook(ruta)["Tabla de estandarización"]

    assert all(celda.data_type == "s" for fila in hoja.iter_rows(min_row=2) for celda in fila if celda.value is not None)


def test_ida_y_vuelta_conserva_las_decisiones(ruta):
    _escribir_celda(ruta, "Tabla de estandarización", "Valor original", "44317", "Decisión", "MANTENER")
    _escribir_celda(ruta, "Tabla de estandarización", "Valor original", "0012", "Decisión", "ACEPTAR")

    decisiones = rx.read_review(ruta).decisions

    assert set(decisiones.table) == {
        ("TIPO DE EXPEDIENTE", "44317", "MANTENER"), ("IDENTIFICACIÓN", "0012", "ACEPTAR"),
        (NAMES, "Ana Ruiz\nLuis Mora", "Ana Ruiz | Luis Mora"), ("SOLICITUD", "Rewquerimiento", "MANTENER")}


def test_las_decisiones_cargadas_que_no_quedaron_en_la_tabla_se_copian(tmp_path):
    cargadas = p.Decisions(table=[("SOLICITUD", "Valor de otro corte", "X"), ("SOLICITUD", "Rewquerimiento", "MANTENER")],
                           catalog=[("SOLICITUD", "VALOR QUE YA NO ESTÁ", "Y")])
    destino = Path(tmp_path) / "r.xlsx"
    rx.write_review(destino, _data(decisions=cargadas))

    leidas = rx.read_review(destino).decisions

    assert ("SOLICITUD", "Valor de otro corte", "X") in leidas.table
    assert leidas.table.count(("SOLICITUD", "Rewquerimiento", "MANTENER")) == 1
    assert ("SOLICITUD", "VALOR QUE YA NO ESTÁ", "Y") in leidas.catalog


def test_ubicaciones_hasta_50(tmp_path):
    entradas = [E("SOLICITUD", "x", "X", None, "error de tipeo", p.AUTOMATIC, list(range(60)))]
    ubicaciones = {i: [f"Reparto_Activo fila {i + 2}"] for i in range(60)}
    destino = Path(tmp_path) / "r.xlsx"
    rx.write_review(destino, _data(entries=entradas, locations=ubicaciones))

    ubicacion = _hoja(destino, "Tabla de estandarización").loc[0, "Ubicación"]

    assert ubicacion.split("; ")[:2] == ["Reparto_Activo fila 2", "Reparto_Activo fila 3"]
    assert len(ubicacion.split("; ")) == 50


# ---------------------------------------------------------------------
# Catálogos y "Corregir a" (RF-42, RF-58)
# ---------------------------------------------------------------------

def test_hoja_de_catalogo_por_columna_y_su_correccion_se_lee(ruta):
    catalogo = _hoja(ruta, "Catálogo SOLICITUD")
    assert list(catalogo.columns) == ["Valor", "Filas", "Corregir a"]
    assert list(catalogo["Valor"]) == ["DEVOLUCIÓN", "POR DETERMINAR"]

    _escribir_celda(ruta, "Catálogo SOLICITUD", "Valor", "DEVOLUCIÓN", "Corregir a", "Devolución de expediente")

    assert rx.read_review(ruta).decisions.catalog == [("SOLICITUD", "DEVOLUCIÓN", "Devolución de expediente")]


# ---------------------------------------------------------------------
# Resumen (RF-43, RF-48, RF-71, RF-76, RF-77, RF-92, RF-95)
# ---------------------------------------------------------------------

def _resumen(ruta):
    return _hoja(ruta, "Resumen")


def test_resumen_explica_el_color_la_observacion_y_la_columna_decision(ruta):
    texto = " ".join(_resumen(ruta).astype(str).values.ravel())

    assert "relleno" in texto and "Observación estandarización" in texto
    assert "ACEPTAR" in texto and "MANTENER" in texto and "Corregir a" in texto


def test_resumen_informa_decisiones_sin_coincidencia_e_invalidas(ruta):
    resumen = _resumen(ruta)

    fila = resumen[resumen["Concepto"] == "Decisiones sin coincidencia"]
    assert list(fila["Valores"]) == ["2"]
    assert resumen["Concepto"].str.contains("«Digital»").any()


def test_resumen_estados_por_columna(ruta):
    resumen = _resumen(ruta)
    filas = resumen[(resumen["Sección"] == "Estados por columna") & (resumen["Columna"] == "DECISIÓN")]

    assert dict(zip(filas["Concepto"], filas["Filas"])) == {p.SUMMARY_UNCHANGED: "6", p.REVIEW: "5",
                                                           "de otra categoría": "0"}


def test_resumen_verificacion_de_totales_con_los_cinco_resultados_y_la_suma(ruta):
    resumen = _resumen(ruta)
    filas = resumen[(resumen["Sección"] == "Verificación de totales") & (resumen["Columna"] == "DECISIÓN")]

    conceptos = dict(zip(filas["Concepto"], filas["Filas"]))
    assert list(conceptos) == [p.EMPTY_TO_UNDETERMINED, p.PENDING_RESULT, p.CEDULA_AS_IS, p.REPLACED,
                               p.UNCHANGED, "Suma"]
    assert conceptos["Suma"] == "11" and conceptos[p.CEDULA_AS_IS] == "0"


def test_resumen_conteos_de_separacion_incluidos_los_que_no_se_separan(ruta):
    resumen = _resumen(ruta)
    filas = resumen[resumen["Sección"] == "Separación de comparecientes"]

    assert dict(zip(filas["Concepto"], filas["Valores"])) == {
        "SEPARADA": "3", "SEPARADA SIN CONFIRMAR": "2", "SIN SEPARAR": "1", "No necesitaban separarse": "10",
        "Filas del resultado": "25"}


def test_resumen_nombre_con_el_que_se_guardo(ruta):
    resumen = _resumen(ruta)

    assert list(resumen.loc[resumen["Concepto"] == "Guardado como", "Valores"]) == ["resultado_revision.xlsx"]


def test_conteos_de_separacion_desde_el_resultado():
    df = pd.DataFrame({"ID REGISTRO": [1, 2, 2, 3, 4, 4, 4],
                       "SEPARACIÓN": ["", "SEPARADA", "SEPARADA", "SIN SEPARAR", "SEPARADA SIN CONFIRMAR"] + ["SEPARADA SIN CONFIRMAR"] * 2})

    assert rx.separation_counts(df) == {"SEPARADA": 1, "SEPARADA SIN CONFIRMAR": 1, "SIN SEPARAR": 1, "": 1,
                                        "filas": 7}


# ---------------------------------------------------------------------
# Leer: RF-47 y la tabla revisada copiada (RF-96, RF-89)
# ---------------------------------------------------------------------

def test_un_excel_cualquiera_da_el_error_de_rf47(tmp_path):
    destino = Path(tmp_path) / "otro.xlsx"
    pd.DataFrame({"A": [1]}).to_excel(destino, index=False)

    with pytest.raises(rx.ReviewFileError):
        rx.read_review(destino)


def test_un_archivo_danado_da_el_error_de_rf47_sin_texto_tecnico(tmp_path):
    # T16: calamine lanza su propia excepción (no OSError ni ValueError).
    destino = Path(tmp_path) / "danado.xlsx"
    destino.write_bytes(b"esto no es un Excel")

    with pytest.raises(rx.ReviewFileError, match="dañado") as error:
        rx.read_review(destino)
    assert "Errno" not in str(error.value)


def test_con_tabla_revisada_el_excel_trae_mapeo_y_al_recargarlo_aplica_la_tabla_y_una_decision_nueva(tmp_path):
    tabla = p.ReviewedTable.from_rows([("SOLICITUD", "Devolucion", "DEVOLUCIÓN"), ("SOLICITUD", "Otro", "OTRO")])
    df = pd.DataFrame({"SOLICITUD": ["Devolucion", "Otro", "Nuevo valor"]})
    primera = p.standardize_parameters(df, table=tabla)
    destino = Path(tmp_path) / "r.xlsx"
    rx.write_review(destino, rx.ReviewData(primera.entries, primera.summary, primera.totals, primera.catalogs,
                                           table=tabla))
    _escribir_celda(destino, "Tabla de estandarización", "Valor original", "Nuevo valor", "Decisión", "VALOR NUEVO")

    archivo = rx.read_review(destino)
    segunda = p.standardize_parameters(df, table=archivo.table, decisions=archivo.decisions)

    assert "Mapeo" in pd.ExcelFile(destino, engine="calamine").sheet_names
    assert list(segunda.df["SOLICITUD"]) == ["DEVOLUCIÓN", "OTRO", "VALOR NUEVO"]
