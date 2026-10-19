"""T20: leer el estilo del archivo original desde el XML del .xlsx, sin cargar las filas. Datos sintéticos."""
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.table import Table, TableStyleInfo

from src import exportation as ex

TIPO = "TIPO DE EXPEDIENTE\n(ELECTRONICO - HIBRIDO)"
ENCABEZADOS = ["No", "RADICADO IUS", TIPO, "NOMBRES_APELLIDOS", "SALA", "SOLICITUD", "DECISIÓN"]


def _original(ruta, rango=False, tabla=True, alto=40.5):
    """Un libro como el de reparto: otra hoja antes, encabezado con salto de línea, anchos, alto y una tabla."""
    libro = openpyxl.Workbook()
    libro.active.title = "Conteo"
    hoja = libro.create_sheet("Reparto")
    hoja.append(ENCABEZADOS)
    for i in range(3):
        hoja.append([i + 1, f"E-2020-{i}", "FÍSICO", f"Persona {i}", "SRVR", "DEVOLUCIÓN", "AGENDADO"])
    for letra, ancho in zip("ABCD", (9.14, 16.14, 22.71, 39.0)):
        hoja.column_dimensions[letra].width = ancho
    if rango:                                   # E:G con un solo <col min="5" max="7">
        hoja.column_dimensions["E"] = ColumnDimension(hoja, index="E", min=5, max=7, width=11.43)
    else:
        for letra, ancho in zip("EFG", (11.43, 25.0, 35.0)):
            hoja.column_dimensions[letra].width = ancho
    if alto:
        hoja.row_dimensions[1].height = alto
    if tabla:
        t = Table(displayName="Tabla12", ref="A1:G4")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        hoja.add_table(t)
    libro.save(ruta)
    return ruta


def _claves(estilo):
    return {k: estilo[k] for k in ("row_height", "anchos_por_nombre", "ancho_default", "table_style", "banded_rows")}


def test_da_lo_mismo_que_el_lector_de_openpyxl(tmp_path):
    ruta = _original(tmp_path / "o.xlsx")

    assert _claves(ex.read_table_layout(ruta, "Reparto")) == _claves(ex.read_header_style(ruta, "Reparto"))


def test_lee_alto_estilo_bandas_y_anchos_por_nombre(tmp_path):
    estilo = ex.read_table_layout(_original(tmp_path / "o.xlsx"), "Reparto")

    assert estilo["row_height"] == 40.5
    assert (estilo["table_style"], estilo["banded_rows"]) == ("TableStyleMedium2", False)
    assert estilo["anchos_por_nombre"][TIPO] == 22.71          # encabezado con salto de línea
    assert estilo["anchos_por_nombre"]["DECISIÓN"] == 35.0


def test_las_columnas_dentro_de_un_rango_toman_su_ancho(tmp_path):
    # openpyxl solo guarda el ancho en la primera columna del rango; las demás
    # caían al ancho promedio aunque en el original miden lo mismo.
    estilo = ex.read_table_layout(_original(tmp_path / "o.xlsx", rango=True), "Reparto")

    assert [estilo["anchos_por_nombre"][c] for c in ("SALA", "SOLICITUD", "DECISIÓN")] == [11.43, 11.43, 11.43]


def test_sin_alto_en_la_primera_fila(tmp_path):
    estilo = ex.read_table_layout(_original(tmp_path / "o.xlsx", alto=None), "Reparto")

    assert estilo["row_height"] is None


def test_sin_tabla_usa_el_lector_de_openpyxl(tmp_path):
    ruta = _original(tmp_path / "o.xlsx", tabla=False)

    estilo = ex.read_table_layout(ruta, "Reparto")

    assert estilo["table_style"] is None
    assert _claves(estilo) == _claves(ex.read_header_style(ruta, "Reparto"))


def test_decodifica_los_saltos_de_linea_escapados_como_los_escribe_excel():
    assert ex._decode_xstring("TIPO DE EXPEDIENTE_x000a_(ELECTRONICO - HIBRIDO)") == TIPO


def test_la_exportacion_ya_no_carga_el_original_con_openpyxl(tmp_path, monkeypatch):
    ruta = _original(tmp_path / "o.xlsx")

    def lento(*a, **k):
        raise AssertionError("no debería abrir el original completo")
    monkeypatch.setattr(ex, "read_header_style", lento)
    salida = Path(tmp_path) / "r.xlsx"

    ex.export_multi_sheet_excel({"Reparto_Activo": pd.DataFrame({"No": [1], TIPO: ["FÍSICO"]})}, salida,
                                style_reference_path=ruta, style_reference_sheet="Reparto", raise_errors=True)

    hoja = openpyxl.load_workbook(salida)["Reparto_Activo"]
    # xlsxwriter guarda el ancho con su propio margen (~0,71 de carácter), como con el lector viejo.
    assert hoja.column_dimensions["B"].width == pytest.approx(22.71 + 0.71, abs=0.05)
    assert hoja.row_dimensions[1].height == 40.5
    assert list(hoja.tables.values())[0].tableStyleInfo.name == "TableStyleMedium2"
