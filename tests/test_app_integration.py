"""
Integración de extremo a extremo: simula lo que dispararían los botones
reales de la GUI (seleccionar -> leer/diagnosticar -> procesar -> se
abre el diálogo de guardado automáticamente, RF-8 -> exportar), pero
llamando directamente al mismo código que usan los callbacks de
gui/app.py en vez de clics reales — no hay una herramienta de
automatización de apps nativas de Windows disponible en este entorno.

IMPORTANTE: esto corre en el escritorio Windows real del usuario (no un
sandbox aislado). _handle_process_done dispara _prompt_save()
AUTOMÁTICAMENTE al despachar el evento "done" (así implementa RF-8), así
que TODAS las funciones de diálogo (filedialog.asksaveasfilename,
messagebox.askyesno/showinfo) se neutralizan ANTES de despachar
cualquier evento — de lo contrario aparece un diálogo real en pantalla
(pasó una vez durante el desarrollo de este test; ver conversación).
También se neutraliza _start_thread para que el export ocurra síncrono
en vez de en un hilo real (no hay mainloop() en la prueba).

No reemplaza una pasada manual real con la app corriendo (ver T8/T9 en
tasks.md), pero sí prueba que el cableado completo entre gui/app.py,
gui/controller.py y src/ funciona con un archivo realista.
"""
import openpyxl
import pandas as pd
from tkinter import messagebox

from gui.app import App
from gui import controller


def _crear_excel_prueba(ruta):
    df = pd.DataFrame({
        "No": [1, 2, 3, 4, 5, 6],
        "RADICADO IUS": ["E-2020-100001", "E-2020-100001", "E-2020-100002", "E-2020-100003", "E-2020-100004",
                         "E-2020-100005"],
        "IDENTIFICACIÓN": ["1000000001", "1000000001", "1000000002\n1000000003", "1000000004", "1000000004",
                           "1000000005\n1000000006"],
        "NOMBRES_APELLIDOS": [
            "  juan pérez gómez  ",
            "JUAN PEREZ GOMEZ",
            "maria lopez\nana torres",
            "Carlos Ruiz",
            "carlos  ruiz",
            "Rosa Marin Tellez Y Hugo Cano Villa",
        ],
        "FECHA": ["01/03/2021", "02/03/2021", "03/03/2021", "04/03/2021", "05/03/2021", "06/03/2021"],
        "REPARTO": ["Marzo 2021"] * 6,
        "CLASIFICACIÓN DEL RADICADO": ["ACTIVO", "ARCHIVADO", "ACTIVO", "ARCHIVADO", "SIM", "ACTIVO"],
        "FUNCIONARIO A CARGO": ["JUAN PEREZ", "FUNCIONARIOS RETIRADOS", "MARIA LOPEZ", "ANA TORRES", "JUAN PEREZ",
                                "JUAN PEREZ"],
        "MAGISTRADO": ["Pedro Suarez", "Pedro Suarez", "Ana Torres", "Ana Torres", "Pedro Suarez", "Pedro Suarez"],
    })
    df.to_excel(ruta, sheet_name="Reparto", index=False)


def test_flujo_completo_seleccion_diagnostico_procesar_guardar(tmp_path, monkeypatch, root):
    # `root` (un tk.Toplevel) viene de tests/conftest.py.
    entrada = tmp_path / "base_prueba.xlsx"
    _crear_excel_prueba(entrada)
    salida = tmp_path / "resultado.xlsx"

    app = App(root)

    # Neutraliza TODO diálogo real y el threading ANTES de despachar
    # ningún evento — _handle_process_done llama _prompt_save() solo,
    # sin que el test lo pida explícitamente (así es RF-8). Desde la spec
    # 002 también showwarning: el aviso de pendientes al guardar (RF-50).
    import gui.app as app_module
    avisos = []
    monkeypatch.setattr(app_module.filedialog, "asksaveasfilename", lambda **k: str(salida))
    monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(messagebox, "showinfo", lambda *a, **k: None)
    monkeypatch.setattr(messagebox, "showerror", lambda *a, **k: None)
    monkeypatch.setattr(messagebox, "showwarning", lambda titulo, texto, **k: avisos.append(texto))
    monkeypatch.setattr(app, "_start_thread", lambda target, *args: target(*args))

    # --- "Seleccionar archivo" (RF-1/RF-2) ---
    app.selected_path = str(entrada)
    eventos = []
    controller.read_and_validate(str(entrada), eventos.append)
    for e in eventos:
        app._dispatch(e)

    assert app.df is not None, "El archivo de prueba debería leerse y validarse"
    assert str(app.process_button["state"]) == "normal"

    # --- "Procesar" (RF-5/RF-6/RF-7) -> dispara guardado automático (RF-8) ---
    eventos = []
    controller.process(app.df, app.source_path, eventos.append)
    for e in eventos:
        app._dispatch(e)

    assert app.process_result is not None
    counts = app.process_result.counts
    # SIM y Multiples_IUS son copias/extracciones que se solapan con
    # Reparto_Activo (ver src/diagnostic.py) — la partición exclusiva
    # real es Activos + Archivados + Retirados. Desde la spec 002 (T16)
    # el controlador usa src/pipeline.py: la fila con 2 comparecientes
    # pasa a 2 filas, y "Carlos Ruiz" / "carlos  ruiz" quedan con el mismo
    # nombre estandarizado y dos IUS distintos. La fila 6 (T17) tiene dos
    # nombres en una línea con dos cédulas: queda SIN SEPARAR, pendiente.
    assert counts["Reparto_Activo"] + counts["Archivados"] + counts["Funcionarios_Retirados"] == 7
    assert counts["SIM"] == 1
    assert counts["Multiples_IUS"] == 2
    assert "Estados por columna" in app.standardization_text.get("1.0", "end")     # RF-43
    assert len(avisos) == 1 and "1 valor en 1 fila" in avisos[0]                    # RF-50

    # _run_export (disparado por el guardado automático de RF-8) puso
    # sus propios eventos en app.queue, no en la lista `eventos` de
    # arriba — se drenan para que _handle_export_done marque saved=True.
    while not app.queue.empty():
        app._dispatch(app.queue.get_nowait())

    revision = tmp_path / "resultado_revision.xlsx"
    assert salida.exists() and revision.exists(), "se guardan el resultado y el Excel de revisión (RF-36)"
    wb = openpyxl.load_workbook(salida)
    assert set(wb.sheetnames) == {
        "Reparto_Activo", "Archivados", "Funcionarios_Retirados", "SIM", "Multiples_IUS",
    }
    hoja = wb["Reparto_Activo"]
    encabezados = [c.value for c in hoja[1]]
    pintadas = [fila[encabezados.index("NOMBRES_APELLIDOS")].value for fila in hoja.iter_rows(min_row=2)
                if fila[encabezados.index("NOMBRES_APELLIDOS")].fill.patternType == "solid"]
    assert pintadas == ["ROSA MARIN TELLEZ Y HUGO CANO VILLA"]                          # RF-40
    tabla = pd.read_excel(revision, sheet_name="Tabla de estandarización", dtype=str)
    pendientes = tabla[tabla["Estado"] == "PARA REVISAR"]
    assert list(pendientes["Valor final o sugerido"]) == ["ROSA MARIN TELLEZ | HUGO CANO VILLA"]
    assert revision.name in app.summary_label["text"]
    assert app.saved is True


def _escribir_decisiones(revision):
    """Lo que haría el funcionario en Excel: ACEPTAR la sugerencia del nombre pegado y una decisión nueva."""
    libro = openpyxl.load_workbook(revision)
    hoja = libro["Tabla de estandarización"]
    encabezados = [c.value for c in hoja[1]]
    col_original, col_decision = encabezados.index("Valor original") + 1, encabezados.index("Decisión") + 1
    for fila in range(2, hoja.max_row + 1):
        if hoja.cell(fila, col_original).value == "Rosa Marin Tellez Y Hugo Cano Villa":
            hoja.cell(fila, col_decision).value = "aceptar"          # RF-72: sin importar mayúsculas
    nueva = hoja.max_row + 1
    hoja.cell(nueva, encabezados.index("Columna") + 1).value = "MAGISTRADO"
    hoja.cell(nueva, col_original).value = "Pedro Suarez"
    hoja.cell(nueva, col_decision).value = "Pedro Suarez Rojas"
    libro.save(revision)


def test_flujo_cargar_revision_aplica_las_decisiones(tmp_path, monkeypatch, root):
    # Spec 002, T18 (RF-44 a RF-46): Procesar -> guardar -> escribir decisiones
    # en el Excel de revisión -> "Cargar revisión" -> el nuevo resultado las refleja.
    entrada = tmp_path / "base_prueba.xlsx"
    _crear_excel_prueba(entrada)
    salidas = iter([tmp_path / "resultado.xlsx", tmp_path / "resultado_2.xlsx"])
    revision = tmp_path / "resultado_revision.xlsx"

    app = App(root)
    import gui.app as app_module
    avisos = []
    monkeypatch.setattr(app_module.filedialog, "asksaveasfilename", lambda **k: str(next(salidas)))
    monkeypatch.setattr(app_module.filedialog, "askopenfilename", lambda **k: str(revision))
    monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(messagebox, "showinfo", lambda *a, **k: None)
    monkeypatch.setattr(messagebox, "showerror", lambda *a, **k: None)
    monkeypatch.setattr(messagebox, "showwarning", lambda titulo, texto, **k: avisos.append(texto))
    monkeypatch.setattr(app, "_start_thread", lambda target, *args: target(*args))

    def drenar():
        while not app.queue.empty():
            app._dispatch(app.queue.get_nowait())

    # Seleccionar, procesar y guardar (como el test de arriba).
    app.selected_path = str(entrada)
    eventos = []
    controller.read_and_validate(str(entrada), eventos.append)
    for e in eventos:
        app._dispatch(e)
    app._on_process_click()
    drenar()
    assert revision.exists() and len(avisos) == 1

    # El funcionario decide en el Excel de revisión y lo carga.
    _escribir_decisiones(revision)
    app._on_load_review_click()
    drenar()

    resultado = tmp_path / "resultado_2.xlsx"
    assert resultado.exists() and (tmp_path / "resultado_2_revision.xlsx").exists()
    activos = pd.read_excel(resultado, sheet_name="Reparto_Activo")
    assert {"ROSA MARIN TELLEZ", "HUGO CANO VILLA"} <= set(activos["NOMBRES_APELLIDOS"])       # ACEPTAR separó
    assert "ROSA MARIN TELLEZ Y HUGO CANO VILLA" not in set(activos["NOMBRES_APELLIDOS"])
    assert "PEDRO SUAREZ ROJAS" in set(activos["MAGISTRADO"])                                # decisión nueva
    assert len(avisos) == 1                                  # ya no quedan pendientes: no hay segundo aviso
    assert not activos["Observación estandarización"].fillna("").str.contains("pendiente").any()
    hoja = openpyxl.load_workbook(resultado)["Reparto_Activo"]
    assert not [c for fila in hoja.iter_rows(min_row=2) for c in fila if c.fill.patternType == "solid"]
    tabla = pd.read_excel(tmp_path / "resultado_2_revision.xlsx", sheet_name="Tabla de estandarización", dtype=str)
    decisiones = dict(zip(tabla["Valor original"], tabla["Decisión"].fillna("")))
    assert decisiones["Rosa Marin Tellez Y Hugo Cano Villa"] == "aceptar"                     # RF-46: se copian
    assert decisiones["Pedro Suarez"] == "Pedro Suarez Rojas"
    assert app.standardization_text.get("1.0", "end").startswith(f"Revisión aplicada: {revision.name}")
    assert app.saved is True
