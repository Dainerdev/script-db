"""T14: ubicaciones, color de pendientes y guardado en par (RF-36, RF-37, RF-40, RF-71, RF-80). Datos sintéticos."""
import datetime
import os
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src import review_excel as rx
from src.diagnostic import split_by_status
from src.exportation import ExportCancelled, export_multi_sheet_excel

TIPO = "TIPO DE EXPEDIENTE\n(ELECTRONICO - HIBRIDO)"


def _df():
    return pd.DataFrame({
        "No": [1, 2, 3, 4],
        "SOLICITUD": ["DEVOLUCIÓN", "agendada", "DEVOLUCIÓN", "x"],
        TIPO: ["FÍSICO", "FÍSICO", "Digital", "FÍSICO"],
        "FECHA": [datetime.datetime(2024, 3, 29)] * 4,
        "_row_id": [10, 11, 12, 13],
        "_pending": [(), ("SOLICITUD",), ("TIPO DE EXPEDIENTE",), ()],
    })


def _relleno(ruta, hoja, fila, columna):
    ws = openpyxl.load_workbook(ruta)[hoja]
    encabezados = [celda.value for celda in ws[1]]
    celda = ws.cell(fila, encabezados.index(columna) + 1)
    return celda.fill.patternType, celda.fill.fgColor.rgb, celda.value


# ---------------------------------------------------------------------
# Color de pendientes en export_multi_sheet_excel (RF-40, Decisión #7)
# ---------------------------------------------------------------------

def test_la_celda_pendiente_tiene_relleno_y_las_demas_no(tmp_path):
    ruta = Path(tmp_path) / "r.xlsx"
    datos = rx.public_frame(_df())

    export_multi_sheet_excel({"Reparto_Activo": datos}, ruta, raise_errors=True,
                             highlight_cells={"Reparto_Activo": [(1, "SOLICITUD")]})

    tipo, color, valor = _relleno(ruta, "Reparto_Activo", 3, "SOLICITUD")
    assert tipo == "solid" and color.endswith(rx.PENDING_FILL[1:]) and valor == "agendada"
    assert _relleno(ruta, "Reparto_Activo", 2, "SOLICITUD")[0] is None
    assert _relleno(ruta, "Reparto_Activo", 3, "No")[0] is None


def test_el_relleno_conserva_fechas_y_vacios(tmp_path):
    ruta = Path(tmp_path) / "r.xlsx"
    datos = rx.public_frame(_df())
    datos.loc[0, "SOLICITUD"] = None

    export_multi_sheet_excel({"H": datos}, ruta, raise_errors=True,
                             highlight_cells={"H": [(0, "SOLICITUD"), (0, "FECHA")]})

    ws = openpyxl.load_workbook(ruta)["H"]
    fecha = ws.cell(2, 4)
    assert fecha.value == datetime.datetime(2024, 3, 29) and fecha.number_format.lower() == "dd/mm/yyyy"
    assert fecha.fill.patternType == "solid"
    assert ws.cell(2, 2).value is None and ws.cell(2, 2).fill.patternType == "solid"


def test_con_estilo_de_tabla_del_original_el_relleno_tambien_queda(tmp_path):
    original = Path(tmp_path) / "original.xlsx"
    with pd.ExcelWriter(original, engine="xlsxwriter") as w:
        pd.DataFrame({"No": [1], "SOLICITUD": ["a"]}).to_excel(w, sheet_name="Reparto", index=False, header=False,
                                                              startrow=1)
        w.sheets["Reparto"].add_table(0, 0, 1, 1, {"style": "Table Style Medium 2",
                                                  "columns": [{"header": "No"}, {"header": "SOLICITUD"}]})
    ruta = Path(tmp_path) / "r.xlsx"

    export_multi_sheet_excel({"Reparto_Activo": rx.public_frame(_df())}, ruta, style_reference_path=original,
                             style_reference_sheet="Reparto", raise_errors=True,
                             highlight_cells={"Reparto_Activo": [(1, "SOLICITUD")]})

    assert _relleno(ruta, "Reparto_Activo", 3, "SOLICITUD")[0] == "solid"
    assert openpyxl.load_workbook(ruta)["Reparto_Activo"].tables


def test_sin_highlight_cells_no_hay_ningun_relleno(tmp_path):
    ruta = Path(tmp_path) / "r.xlsx"
    datos = rx.public_frame(_df())

    export_multi_sheet_excel({"H": datos}, ruta, raise_errors=True)

    ws = openpyxl.load_workbook(ruta)["H"]
    assert all(celda.fill.patternType is None for fila in ws.iter_rows() for celda in fila)
    assert [celda.value for celda in ws[1]] == list(datos.columns)


# ---------------------------------------------------------------------
# Columnas internas, pendientes a pintar y ubicaciones (RF-37)
# ---------------------------------------------------------------------

def test_las_columnas_internas_no_salen_en_el_resultado():
    datos = _df()
    datos["Observación estandarización"] = [""] * 4
    datos["_marks"] = [()] * 4

    assert list(rx.public_frame(datos).columns) == ["No", "SOLICITUD", TIPO, "FECHA", "Observación estandarización"]


def test_celdas_a_pintar_salen_de_pending_con_el_encabezado_real():
    pintar = rx.highlight_cells({"Reparto_Activo": _df()})

    assert pintar == {"Reparto_Activo": [(1, "SOLICITUD"), (2, TIPO)]}


def test_la_ubicacion_apunta_a_la_fila_correcta_despues_de_separar_por_estado():
    df = pd.DataFrame({
        "CLASIFICACIÓN DEL RADICADO": ["NUEVO", "ARCHIVADO", "NUEVO", "ARCHIVADO"],
        "FUNCIONARIO A CARGO": ["ANA", "LUIS", "ANA", "LUIS"],
        "_row_id": [0, 1, 2, 3],
    })
    grupos = split_by_status(df)
    hojas = {"Reparto_Activo": grupos["activos"], "Archivados": grupos["archivados"]}

    ubicaciones = rx.locations_by_row(hojas)

    assert ubicaciones[2] == ["Reparto_Activo fila 3"]
    assert ubicaciones[3] == ["Archivados fila 3"]


def test_una_persona_separada_tiene_una_ubicacion_por_fila():
    hojas = {"Reparto_Activo": pd.DataFrame({"_row_id": [7, 7, 8]})}

    assert rx.locations_by_row(hojas)[7] == ["Reparto_Activo fila 2", "Reparto_Activo fila 3"]


# ---------------------------------------------------------------------
# Guardado en par (RF-36, RF-71, RF-80, Decisión #8)
# ---------------------------------------------------------------------

def _escribe(texto):
    return lambda ruta: Path(ruta).write_text(texto, encoding="utf-8")


def test_guarda_el_resultado_y_la_revision_junto(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"

    r, v = rx.save_pair(resultado, _escribe("resultado"), _escribe("revision"))

    assert (r, v) == (resultado, Path(tmp_path) / "salida_revision.xlsx")
    assert r.read_text(encoding="utf-8") == "resultado" and v.read_text(encoding="utf-8") == "revision"
    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["salida.xlsx", "salida_revision.xlsx"]


def test_si_ya_existe_la_revision_guarda_con_numero_y_no_la_toca(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"
    (Path(tmp_path) / "salida_revision.xlsx").write_text("vieja", encoding="utf-8")
    (Path(tmp_path) / "salida_revision (2).xlsx").write_text("vieja 2", encoding="utf-8")

    _, v = rx.save_pair(resultado, _escribe("resultado"), _escribe("nueva"))

    assert v.name == "salida_revision (3).xlsx" and v.read_text(encoding="utf-8") == "nueva"
    assert (Path(tmp_path) / "salida_revision.xlsx").read_text(encoding="utf-8") == "vieja"


def test_un_fallo_al_escribir_la_revision_no_deja_ninguno(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"

    def falla(ruta):
        Path(ruta).write_text("a medias", encoding="utf-8")
        raise OSError("disco lleno")

    with pytest.raises(OSError):
        rx.save_pair(resultado, _escribe("resultado"), falla)

    assert list(Path(tmp_path).iterdir()) == []


def test_un_fallo_no_borra_el_resultado_que_ya_existia(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"
    resultado.write_text("anterior", encoding="utf-8")

    with pytest.raises(OSError):
        rx.save_pair(resultado, _escribe("nuevo"), lambda ruta: (_ for _ in ()).throw(OSError("x")))

    assert resultado.read_text(encoding="utf-8") == "anterior"
    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["salida.xlsx"]


def test_si_falla_el_ultimo_reemplazo_se_deshace_el_primero(tmp_path, monkeypatch):
    resultado = Path(tmp_path) / "salida.xlsx"
    resultado.write_text("anterior", encoding="utf-8")
    real = os.replace

    def replace(origen, destino):
        if str(destino).endswith("_revision.xlsx"):
            raise PermissionError("abierto en Excel")
        return real(origen, destino)

    monkeypatch.setattr(rx.os, "replace", replace)

    with pytest.raises(PermissionError):
        rx.save_pair(resultado, _escribe("nuevo"), _escribe("revision"))

    assert resultado.read_text(encoding="utf-8") == "anterior"
    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["salida.xlsx"]


def test_cancelar_no_deja_ninguno(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"

    def cancela(ruta):
        raise ExportCancelled("cancelado")

    with pytest.raises(ExportCancelled):
        rx.save_pair(resultado, cancela, _escribe("revision"))

    assert list(Path(tmp_path).iterdir()) == []


# T17: un resultado sin Excel de revisión (el de la spec 001) también se
# guarda por un temporal, así cancelar nunca deja un archivo a medias (RF-15).

def test_sin_revision_guarda_solo_el_resultado(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"

    r, v = rx.save_pair(resultado, _escribe("resultado"), None)

    assert (r, v) == (resultado, None)
    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["salida.xlsx"]


def test_sin_revision_cancelar_conserva_el_resultado_que_ya_existia(tmp_path):
    resultado = Path(tmp_path) / "salida.xlsx"
    resultado.write_text("anterior", encoding="utf-8")

    def cancela(ruta):
        Path(ruta).write_text("a medias", encoding="utf-8")
        raise ExportCancelled("cancelado")

    with pytest.raises(ExportCancelled):
        rx.save_pair(resultado, cancela, None)

    assert resultado.read_text(encoding="utf-8") == "anterior"
    assert sorted(p.name for p in Path(tmp_path).iterdir()) == ["salida.xlsx"]
