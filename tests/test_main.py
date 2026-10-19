"""T19: main.py usa el pipeline de la spec 002 (RF-60, RF-61). Datos sintéticos."""
import queue
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

import main
from gui import controller
from gui.app import App

PEGADO = "Rosa Marin Tellez Y Hugo Cano Villa"


def _crear_entrada(ruta):
    pd.DataFrame({
        "No": [1, 2, 3, 4, 5],
        "RADICADO IUS": [f"E-2020-10000{i}" for i in range(5)],
        "IDENTIFICACIÓN": ["1000000001", "1000000002\n1000000003", "1000000004\n1000000005", "1000000006", "Caso 3"],
        "NOMBRES_APELLIDOS": ["  juan pérez gómez ", "Maria Lopez Rios\nAna Torres Paz", PEGADO, "Carlos Ruiz Vega",
                              "Solicitud de informe"],
        "FECHA": ["01/03/2021"] * 5,
        "REPARTO": ["Marzo 2021"] * 5,
        "CLASIFICACIÓN DEL RADICADO": ["ACTIVO", "ACTIVO", "ARCHIVADO", "SIM", "ACTIVO"],
        "FUNCIONARIO A CARGO": ["JUAN PEREZ", "MARIA LOPEZ", "JUAN PEREZ", "FUNCIONARIOS RETIRADOS", "MARIA LOPEZ"],
        "MAGISTRADO": ["Pedro Suarez", "Pedro Suárez", "Ana Torres", "Ana Torres", "Pedro Suarez"],
        "SOLICITUD": ["Devolución", "devolucion", "REQUERIMIENTO", "Devoución", None],
    }).to_excel(ruta, sheet_name="Reparto", index=False)


@pytest.fixture
def config(tmp_path, monkeypatch):
    """Apunta las rutas de configuración de main.py a archivos sintéticos en tmp_path."""
    entrada = tmp_path / "entrada.xlsx"
    _crear_entrada(entrada)
    (tmp_path / "main").mkdir()
    monkeypatch.setattr(main, "FILE_EXCEL", str(entrada))
    monkeypatch.setattr(main, "FILE_RESULTS", str(tmp_path / "main" / "resultado.xlsx"))
    monkeypatch.setattr(main, "FILE_REVIEW", None)
    return tmp_path


def _hojas(ruta):
    return pd.read_excel(ruta, sheet_name=None, dtype=str)


def test_main_guarda_el_resultado_y_el_excel_de_revision_juntos(config):
    main.main()

    carpeta = config / "main"
    assert sorted(p.name for p in carpeta.iterdir()) == ["resultado.xlsx", "resultado_revision.xlsx"]   # RF-60
    hojas = _hojas(carpeta / "resultado.xlsx")
    assert list(hojas) == ["Reparto_Activo", "Archivados", "Funcionarios_Retirados", "SIM", "Multiples_IUS"]
    assert "Observación estandarización" in hojas["Reparto_Activo"].columns
    assert "Tabla de estandarización" in _hojas(carpeta / "resultado_revision.xlsx")


def test_main_y_la_gui_dan_el_mismo_resultado_y_la_misma_tabla(config):
    main.main()

    # Lo mismo por la GUI: controller.process y el guardado de "Guardar resultado...".
    (config / "gui").mkdir()
    eventos = []
    df = pd.read_excel(main.FILE_EXCEL, sheet_name="Reparto")
    controller.process(df, main.FILE_EXCEL, eventos.append)
    falsa = SimpleNamespace(queue=queue.Queue(), cancel_requested=False)
    App._run_export(falsa, eventos[-1][1], str(config / "gui" / "resultado.xlsx"))

    for nombre in ("resultado.xlsx", "resultado_revision.xlsx"):
        de_main, de_gui = _hojas(config / "main" / nombre), _hojas(config / "gui" / nombre)
        assert list(de_main) == list(de_gui), nombre
        for hoja in de_main:
            pd.testing.assert_frame_equal(de_main[hoja], de_gui[hoja], obj=f"{nombre} · {hoja}")


def test_main_pinta_las_celdas_pendientes(config):
    main.main()

    hoja = openpyxl.load_workbook(config / "main" / "resultado.xlsx")["Archivados"]
    encabezados = [c.value for c in hoja[1]]
    celda = hoja.cell(2, encabezados.index("NOMBRES_APELLIDOS") + 1)
    assert celda.value == PEGADO.upper() and celda.fill.patternType == "solid"         # RF-40


def test_main_con_file_review_aplica_las_decisiones(config, monkeypatch):
    main.main()
    revision = config / "main" / "resultado_revision.xlsx"
    libro = openpyxl.load_workbook(revision)
    hoja = libro["Tabla de estandarización"]
    encabezados = [c.value for c in hoja[1]]
    nueva = hoja.max_row + 1
    hoja.cell(nueva, encabezados.index("Columna") + 1).value = "MAGISTRADO"
    hoja.cell(nueva, encabezados.index("Valor original") + 1).value = "Ana Torres"
    hoja.cell(nueva, encabezados.index("Decisión") + 1).value = "Ana Torres Paz"
    for fila in range(2, hoja.max_row + 1):
        if hoja.cell(fila, encabezados.index("Valor original") + 1).value == PEGADO:
            hoja.cell(fila, encabezados.index("Decisión") + 1).value = "ACEPTAR"
    libro.save(revision)
    monkeypatch.setattr(main, "FILE_REVIEW", str(revision))                                    # RF-61
    monkeypatch.setattr(main, "FILE_RESULTS", str(config / "main" / "resultado_2.xlsx"))

    main.main()

    hojas = _hojas(config / "main" / "resultado_2.xlsx")
    todas = pd.concat([hojas["Reparto_Activo"], hojas["Archivados"], hojas["Funcionarios_Retirados"]])
    assert "ANA TORRES PAZ" in set(todas["MAGISTRADO"])
    assert {"ROSA MARIN TELLEZ", "HUGO CANO VILLA"} <= set(todas["NOMBRES_APELLIDOS"])
    tabla = _hojas(config / "main" / "resultado_2_revision.xlsx")["Tabla de estandarización"]
    assert "Ana Torres Paz" in set(tabla["Decisión"].dropna())                                  # RF-46


def test_main_con_un_file_review_que_no_sirve_no_escribe_nada(config, monkeypatch, capsys):
    otro = config / "otro.xlsx"
    pd.DataFrame({"A": [1]}).to_excel(otro, index=False)
    monkeypatch.setattr(main, "FILE_REVIEW", str(otro))

    main.main()

    assert list((config / "main").iterdir()) == []
    assert "no es un Excel de revisión" in capsys.readouterr().out                              # RF-47


def test_main_sin_archivo_de_entrada_no_falla(config, monkeypatch):
    monkeypatch.setattr(main, "FILE_EXCEL", str(config / "no_existe.xlsx"))

    main.main()

    assert list((config / "main").iterdir()) == []


def test_main_crea_la_carpeta_de_resultados_si_no_existe(config, monkeypatch):
    monkeypatch.setattr(main, "FILE_RESULTS", str(config / "nueva" / "resultado.xlsx"))

    main.main()

    assert (config / "nueva" / "resultado_revision.xlsx").exists()
